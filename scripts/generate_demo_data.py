"""
Synthetic Demo Data Generator for VOLTRA MOBILITY
Generates realistic multi-format synthetic files:
- Suppliers.xlsx
- Inventory.xlsx
- Production_Process.pdf
- Machine_Requirements.pdf
- Supply_Contracts.pdf
- Operations_SOP.docx
Also generates .txt and .csv equivalents for maximum parsing resilience.
"""

import os
import csv
import json


def create_demo_files(output_dir: str):
    os.makedirs(output_dir, exist_ok=True)
    
    # 1. Suppliers Data (CSV & Excel)
    suppliers_data = [
        ["Supplier_Name", "Component_Supplied", "Category", "Lead_Time_Days", "Backup_Supplier", "Backup_SLA_Days", "Contract_ID", "Risk_Level"],
        ["Apex Microelectronics", "Microcontroller MCU-900", "Avionics / Microchips", 21, "None", "N/A", "MSA-4401", "CRITICAL - SINGLE SOURCE"],
        ["VoltCell Energy", "Lithium-Ion Battery Cells", "Energy Storage", 14, "Amperex Dynamics", 3, "MSA-8812", "LOW - BACKUP CERTIFIED"],
        ["StatorTech Precision", "Electric Traction Motor Stator", "Electric Powertrain", 10, "GlobalDrive Dynamics", 14, "MSA-3109", "MODERATE"],
        ["Sensoryx Corp", "Ultrasonic Radar Sensors", "ADAS Sensors", 7, "None", "N/A", "MSA-9021", "MODERATE"]
    ]
    
    # Write CSV
    sup_csv_path = os.path.join(output_dir, "Suppliers.csv")
    with open(sup_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(suppliers_data)
        
    # Write Excel (.xlsx) if openpyxl available
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vendors"
        for r_idx, row in enumerate(suppliers_data, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        wb.save(os.path.join(output_dir, "Suppliers.xlsx"))
    except ImportError:
        pass
        
    # 2. Inventory Data (CSV & Excel)
    inventory_data = [
        ["Component_Name", "Current_Stock", "Daily_Burn_Rate", "Coverage_Days", "Safety_Stock_Threshold", "Warehouse_Location"],
        ["Microcontroller MCU-900", 2500, 500, 5.0, 3500, "Bay A-12 (Climate Controlled)"],
        ["Lithium-Ion Battery Cells", 12000, 1000, 12.0, 5000, "High Voltage Bunker 2"],
        ["Electric Traction Motor Stator", 4000, 400, 10.0, 2000, "Bay C-04"],
        ["Ultrasonic Radar Sensors", 3000, 300, 10.0, 1500, "Bay B-08"]
    ]
    
    inv_csv_path = os.path.join(output_dir, "Inventory.csv")
    with open(inv_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(inventory_data)
        
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock_Master"
        for r_idx, row in enumerate(inventory_data, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        wb.save(os.path.join(output_dir, "Inventory.xlsx"))
    except ImportError:
        pass
        
    # 3. Production Process Text & PDF
    prod_text = """VOLTRA MOBILITY // MANUFACTURING PROCESS SPECIFICATION
Document ID: PROC-2026-ENG-08
Classification: Internal Manufacturing Standard

1. EXECUTIVE OVERVIEW
Voltra Mobility operates a continuous synchronous automotive production line for two commercial EV platforms:
- Voltra Model V-1 (All-electric urban passenger sedan, target 500 units/day)
- Voltra Commercial Fleet Van (Light commercial delivery van, target 150 units/day)

2. ASSEMBLY LINE ARCHITECTURE & DEPENDENCIES
Line 1: Powertrain Integration
- Integrates Electric Traction Motor Stator assemblies and high-voltage battery modules.
- Output: Complete high-voltage chassis powertrain unit. Feeds directly into Line 3.

Line 2: ECU Sub-assembly & Avionics
- Primary Machine: SMT Robot #4 (Surface Mount Technology high-speed placer).
- Required Component: Microcontroller MCU-900 (Supplied exclusively by Apex Microelectronics).
- Machine B (SMT Robot #4) places MCU-900 chipsets onto multilayer motherboard PCBs.
- Output: Vehicle Central Avionics and Electronic Control Unit (ECU Sub-assembly).
- Dependency: Line 2 output feeds directly into Final Vehicle Synthesis on Line 3.

Line 3: Final Vehicle Synthesis
- Merges chassis, Line 1 powertrain assemblies, Line 2 ECU modules, and sensor harnesses.
- Outputs finished commercial products: Voltra Model V-1 and Voltra Commercial Fleet Van.
- CRITICAL PATH: Failure of Line 2 ECU Sub-assembly halts Line 3 final road-readiness certification immediately.
"""
    with open(os.path.join(output_dir, "Production_Process.txt"), "w", encoding="utf-8") as f:
        f.write(prod_text)
        
    # Also write fake/simple PDF stream
    _write_simple_pdf(os.path.join(output_dir, "Production_Process.pdf"), "Production Process Specification", prod_text)

    # 4. Machine Requirements Text & PDF
    mach_text = """VOLTRA MOBILITY // TOOLING & MACHINE SPECIFICATIONS
Document ID: MACH-SPEC-404

Machine B: SMT Robot #4
- Station Location: Assembly Line 2 (ECU Sub-assembly)
- Function: High-speed placement of microchips onto vehicle avionics circuit boards.
- Critical Feed Component: Microcontroller MCU-900 (taped reel format).
- Operating Dependency: SMT Robot #4 cannot substitute alternative microchip architectures without full recalibration and firmware recompilation (estimated 90 days).

Battery Pack Laser Welder
- Station Location: Assembly Line 1 (Powertrain Integration)
- Function: Automated robotic busbar welding for Lithium-Ion Battery Cells.
- Cycle Time: 45 seconds per module.

Stator Winding Automated Press
- Station Location: Assembly Line 1 (Powertrain Integration)
- Function: Press-fits Electric Traction Motor Stator core into electric motor housing.

Automated Test Rig #2
- Station Location: Line 3 Inspection Station
- Function: Calibration of Ultrasonic Radar Sensors for collision avoidance and ADAS.
"""
    with open(os.path.join(output_dir, "Machine_Requirements.txt"), "w", encoding="utf-8") as f:
        f.write(mach_text)
    _write_simple_pdf(os.path.join(output_dir, "Machine_Requirements.pdf"), "Machine Requirements", mach_text)

    # 5. Supply Contracts Text & PDF
    contracts_text = """VOLTRA MOBILITY // MASTER SUPPLY AGREEMENTS SUMMARY
Legal Repository: LEGAL-MSA-2025-ARCHIVE

CONTRACT MSA-4401: APEX MICROELECTRONICS
- Parties: Voltra Mobility Inc. & Apex Microelectronics Corp.
- Component: Microcontroller MCU-900 (Automotive Grade 32-bit SoC).
- Section 3 (Sole Sourcing & Exclusivity): Apex Microelectronics is designated as the EXCLUSIVE supplier for MCU-900. Voltra covenants not to procure compatible silicon from alternative foundries.
- Section 7 (Backup & Contingency): NO certified secondary foundry exists. Requalification of alternative silicon requires automotive safety certification (ISO 26262 ASIL-D) taking a minimum of 90 days.
- SOURCING STATUS: SINGLE POINT OF FAILURE. No documented backup supplier certified.

CONTRACT MSA-8812: VOLTCELL ENERGY
- Parties: Voltra Mobility Inc. & VoltCell Energy Ltd.
- Component: Lithium-Ion Battery Cells (4680 cylindrical format).
- Section 8.2 (Emergency Secondary Sourcing): In the event of supply interruption exceeding 48 hours, Voltra Mobility is pre-authorized to activate approved secondary supplier Amperex Dynamics.
- Activation SLA: Guaranteed initial shipment within 3 business days from Amperex Dynamics domestic facility.
- SOURCING STATUS: PROTECTED BY CERTIFIED BACKUP.
"""
    with open(os.path.join(output_dir, "Supply_Contracts.txt"), "w", encoding="utf-8") as f:
        f.write(contracts_text)
    _write_simple_pdf(os.path.join(output_dir, "Supply_Contracts.pdf"), "Master Supply Contracts", contracts_text)

    # 6. Operations SOP Text & DOCX
    sop_text = """VOLTRA MOBILITY // STANDARD OPERATING PROCEDURE SOP-704
Title: Supply Chain Disruption & Factory Contingency Protocol
Effective Date: January 15, 2026

1. PURPOSE
Defines operational response thresholds when critical component deliveries are halted.

2. INVENTORY RUNWAY EVALUATION
2.1 Buffer Calculation: Days of Inventory = Current Stock / Daily Consumption.
2.2 Microcontroller MCU-900 Stock Policy:
    - Normal Buffer: 5 Days of active stock (2,500 units at 500 units/day consumption).
    - If supplier outage duration exceeds 5 days: Critical Gap occurs.
    - Day 1-5: Production operates normally using on-hand buffer.
    - Day 6+: Total line stoppage of Line 2 (ECU Sub-assembly) and Line 3 (Voltra Model V-1).

3. ALTERNATIVE SOURCING PROTOCOLS
3.1 For components with approved secondary suppliers (e.g. VoltCell -> Amperex Dynamics):
    Initiate secondary purchase order within 12 hours of confirmed primary disruption.
3.2 For single-source components with NO documented alternative (e.g. Apex Microelectronics MCU-900):
    Plant manager must declare Force Majeure operational suspension upon stock depletion.
"""
    with open(os.path.join(output_dir, "Operations_SOP.txt"), "w", encoding="utf-8") as f:
        f.write(sop_text)
        
    # Write DOCX if python-docx available
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("VOLTRA MOBILITY // SOP-704", 0)
        for line in sop_text.splitlines():
            if line.startswith("1.") or line.startswith("2.") or line.startswith("3."):
                doc.add_heading(line, level=1)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(os.path.join(output_dir, "Operations_SOP.docx"))
    except ImportError:
        pass


def _write_simple_pdf(filepath: str, title: str, content: str):
    """Generates a standard compliant minimal PDF file with exact xref byte offsets."""
    lines = content.splitlines()
    stream_lines = [f"BT /F1 12 Tf 50 750 Td ({title}) Tj ET"]
    y = 720
    for l in lines[:32]:
        clean = l.replace("(", "[").replace(")", "]").replace("\\", "/")
        if clean.strip():
            stream_lines.append(f"BT /F1 9 Tf 50 {y} Td ({clean[:85]}) Tj ET")
            y -= 18
            if y < 60:
                break
                
    stream_str = "\n".join(stream_lines)
    stream_bytes = stream_str.encode("latin1", errors="ignore")
    stream_len = len(stream_bytes)
    
    header = b"%PDF-1.4\n"
    obj1 = b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    obj2 = b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
    obj3 = b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    obj4_prefix = f"4 0 obj\n<< /Length {stream_len} >>\nstream\n".encode("latin1")
    obj4_suffix = b"\nendstream\nendobj\n"
    obj5 = b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    
    pos1 = len(header)
    pos2 = pos1 + len(obj1)
    pos3 = pos2 + len(obj2)
    pos4 = pos3 + len(obj3)
    pos5 = pos4 + len(obj4_prefix) + stream_len + len(obj4_suffix)
    xref_pos = pos5 + len(obj5)
    
    xref = f"xref\n0 6\n0000000000 65535 f \n{pos1:010d} 00000 n \n{pos2:010d} 00000 n \n{pos3:010d} 00000 n \n{pos4:010d} 00000 n \n{pos5:010d} 00000 n \n".encode("latin1")
    trailer = f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode("latin1")
    
    with open(filepath, "wb") as f:
        f.write(header)
        f.write(obj1)
        f.write(obj2)
        f.write(obj3)
        f.write(obj4_prefix)
        f.write(stream_bytes)
        f.write(obj4_suffix)
        f.write(obj5)
        f.write(xref)
        f.write(trailer)



if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "data/demo"
    create_demo_files(out)
    print(f"Successfully generated synthetic demo files in {out}")
